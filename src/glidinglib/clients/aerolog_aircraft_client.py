from __future__ import annotations

import json
import shutil
from dataclasses import asdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from glidinglib.mappers.aerolog_aircraft_mapper import map_aerolog_aircraft
from glidinglib.models.aerolog_aircraft_model import AerologAircraft
from glidinglib.utils.app_paths import get_app_data_dir


class AerologAircraftClient:
    DEFAULT_CACHE_FILE = "aerolog_aircraft.json"
    DEFAULT_EXCEL_CACHE_FILE = "aerolog_aircraft.xlsx"

    def __init__(
        self,
        app_name: str = "GlidingLib",
        cache_dir: str | Path | None = None,
        cache_file: str = DEFAULT_CACHE_FILE,
        excel_cache_file: str = DEFAULT_EXCEL_CACHE_FILE,
    ):
        self.cache_dir = Path(cache_dir) if cache_dir else get_app_data_dir(app_name)
        self.cache_file = cache_file
        self.excel_cache_file = excel_cache_file

        self._records: list[AerologAircraft] | None = None

        self._by_registration: dict[str, AerologAircraft] = {}
        self._by_short_registration: dict[str, AerologAircraft] = {}
        self._by_competition_registration: dict[str, AerologAircraft] = {}

    @property
    def cache_path(self) -> Path:
        return self.cache_dir / self.cache_file

    @property
    def excel_cache_path(self) -> Path:
        return self.cache_dir / self.excel_cache_file

    def load(self) -> list[AerologAircraft]:
        """
        Load aircraft from the JSON cache.
        """
        if not self.cache_path.exists():
            raise FileNotFoundError(
                f"No Aerolog aircraft cache found at {self.cache_path}. "
                "Call update_cache_from_excel() first."
            )

        with self.cache_path.open("r", encoding="utf-8") as f:
            rows = json.load(f)

        records = [AerologAircraft(**row) for row in rows]

        self._records = records
        self._build_indexes(records)

        return records

    def update_cache_from_excel(
        self,
        excel_path: str | Path,
        sheet_name: str | None = None,
        table_name: str = "Table1",
    ) -> list[AerologAircraft]:
        """
        Update the cache from an Aerolog aircraft Excel download.

        The Excel file is copied into the app cache directory and a JSON cache
        is written from the contents of Table1.
        """
        excel_path = Path(excel_path)

        if not excel_path.exists():
            raise FileNotFoundError(excel_path)

        self.cache_dir.mkdir(parents=True, exist_ok=True)

        shutil.copy2(excel_path, self.excel_cache_path)

        raw_rows = self._read_excel_table(
            self.excel_cache_path,
            sheet_name=sheet_name,
            table_name=table_name,
        )

        records = [map_aerolog_aircraft(row) for row in raw_rows]

        self._save_json_cache(records)

        self._records = records
        self._build_indexes(records)

        return records

    def _read_excel_table(
        self,
        excel_path: Path,
        sheet_name: str | None,
        table_name: str,
    ) -> list[dict[str, Any]]:
        workbook = load_workbook(excel_path, data_only=True)

        worksheet = workbook[sheet_name] if sheet_name else workbook.active

        table = worksheet.tables.get(table_name)

        if table is None:
            raise ValueError(
                f"Table {table_name!r} not found in worksheet {worksheet.title!r}"
            )

        cells = worksheet[table.ref]

        rows = list(cells)
        if not rows:
            return []

        headers = [
            str(cell.value or "").strip()
            for cell in rows[0]
        ]

        output: list[dict[str, Any]] = []

        for row in rows[1:]:
            values = [cell.value for cell in row]
            record = dict(zip(headers, values))

            if any(value not in (None, "") for value in values):
                output.append(record)

        return output

    def _save_json_cache(self, records: list[AerologAircraft]) -> None:
        self.cache_dir.mkdir(parents=True, exist_ok=True)

        with self.cache_path.open("w", encoding="utf-8") as f:
            json.dump(
                [asdict(record) for record in records],
                f,
                indent=2,
            )

    def records(self) -> list[AerologAircraft]:
        self._ensure_loaded()
        return self._records or []

    def find_by_registration(
        self,
        registration: str,
    ) -> AerologAircraft | None:
        self._ensure_loaded()
        return self._by_registration.get(self._normalise_registration(registration))

    def find_by_short_registration(
        self,
        short_registration: str,
    ) -> AerologAircraft | None:
        self._ensure_loaded()
        return self._by_short_registration.get(
            self._normalise_registration(short_registration)
        )

    def find_by_competition_registration(
        self,
        competition_registration: str,
    ) -> AerologAircraft | None:
        self._ensure_loaded()
        return self._by_competition_registration.get(
            self._normalise_registration(competition_registration)
        )

    def _ensure_loaded(self) -> None:
        if self._records is None:
            self.load()

    def _build_indexes(
        self,
        records: list[AerologAircraft],
    ) -> None:
        self._by_registration.clear()
        self._by_short_registration.clear()
        self._by_competition_registration.clear()

        for record in records:
            registration = self._normalise_registration(record.registration)
            short_registration = self._normalise_registration(record.short_registration)
            competition_registration = self._normalise_registration(
                record.competition_registration
            )

            if registration:
                self._by_registration[registration] = record

            if short_registration:
                self._by_short_registration[short_registration] = record

            if competition_registration:
                self._by_competition_registration[competition_registration] = record

    def _normalise_registration(self, value: str) -> str:
        return (
            str(value or "")
            .upper()
            .replace("-", "")
            .replace(" ", "")
            .strip()
        )